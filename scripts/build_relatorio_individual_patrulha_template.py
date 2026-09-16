"""Convert the partner's real Relatório Individual por patrulha mockup into
a docxtpl (Jinja2) template, following the same tag style as the Relatório
Bimestral template (ICMBio-patrol_analysis's
scripts/build_relatorio_bimestral_template.py) and APN's real, working
patrol/event report template.

Run once (or whenever the partner supplies a revised mockup):
    pixi run -e default python scripts/build_relatorio_individual_patrulha_template.py

What this does, matching the mockup's own layout, read from
/Users/zak/Downloads/Relatório Individual por patrulha_modelo.docx:
  - Appends `{{ participantes }}` after the "Participantes:" line.
  - Fills in the stat table's 6 empty value cells (ID da patrulha, Tipo de
    patrulha, Líder da patrulha, Distância percorrida, Data e hora de
    início, Data e hora do fim) — a plain label/value table, not a loop.
  - Replaces the single embedded map screenshot with `{{ patrol_map }}`.
  - Turns the Eventos table's one blank data row into a docxtpl row-loop
    over `events` (ID do evento / Tipo de evento / Detalhes do evento —
    the 4th, blank header column is left blank in every row, same as the
    Relatório Bimestral template's Ameaças table).
  - Turns the Registros Fotográficos table's one blank data row into a
    row-loop over `photos` (info / image), same 3-row pattern.

Context keys expected at render time (see ecoscope_workflows_ext_icmbio's
prepare_individual_patrol_report_context / generate_individual_patrol_report):
    patrol_id, patrol_type, patrol_leader, participantes, total_distance_km,
    start_time, end_time, events (list of dicts), patrol_map (InlineImage),
    photos (list of {info, image} dicts).
"""

import copy
from pathlib import Path

import docx
from docx.table import Table, _Row
from docx.text.paragraph import Paragraph

SRC = Path("/Users/zak/Downloads/Relatório Individual por patrulha_modelo.docx")
DST = Path(__file__).parent.parent / "resources" / "templates" / "relatorio_individual_patrulha_template.docx"


# ── low-level helpers (same as build_relatorio_bimestral_template.py) ─────────

def find_paragraph(doc: docx.Document, prefix: str) -> Paragraph:
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise ValueError(f"paragraph not found (prefix={prefix!r})")


def append_jinja(paragraph: Paragraph, expr: str) -> None:
    ref = paragraph.runs[0] if paragraph.runs else None
    run = paragraph.add_run(expr)
    if ref is not None:
        run.bold = ref.bold
        run.font.size = ref.font.size
        run.font.name = ref.font.name


def clear_cell(cell) -> Paragraph:
    cell.text = ""
    return cell.paragraphs[0]


def set_cell_jinja(cell, expr: str, bold: bool = False):
    p = clear_cell(cell)
    r = p.add_run(expr)
    r.bold = bold
    return r


def image_paragraphs(doc: docx.Document) -> list[Paragraph]:
    """Body paragraphs containing an embedded drawing, in document order."""
    out = []
    for p in doc.paragraphs:
        if p._p.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}graphic"):
            out.append(p)
    return out


def replace_image_with_jinja(paragraph: Paragraph, expr: str) -> None:
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    paragraph.add_run(expr)


def clone_row(table: Table, ref_row, after: bool) -> _Row:
    new_tr = copy.deepcopy(ref_row._tr)
    if after:
        ref_row._tr.addnext(new_tr)
    else:
        ref_row._tr.addprevious(new_tr)
    return _Row(new_tr, table)


def make_row_loop(table: Table, data_row_index: int, loop_var: str, cell_exprs: list[str]) -> None:
    """Turn table.rows[data_row_index] (a single blank data row) into a
    docxtpl row-loop: a loop-start row, the data row (filled with
    cell_exprs), and a loop-end row."""
    data_row = table.rows[data_row_index]
    ncols = len(data_row.cells)

    loop_start = clone_row(table, data_row, after=False)
    loop_end = clone_row(table, data_row, after=True)

    set_cell_jinja(loop_start.cells[0], "{%tr for item in " + loop_var + " %}")
    for c in loop_start.cells[1:]:
        clear_cell(c)

    for cell, expr in zip(data_row.cells, cell_exprs):
        set_cell_jinja(cell, expr)
    for cell in data_row.cells[len(cell_exprs):]:
        clear_cell(cell)

    set_cell_jinja(loop_end.cells[0], "{%tr endfor %}")
    for c in loop_end.cells[1:]:
        clear_cell(c)

    assert len(data_row.cells) == ncols


# ── main ──────────────────────────────────────────────────────────────────────

def main() -> None:
    doc = docx.Document(str(SRC))

    # ── "Participantes:" line ─────────────────────────────────────────────────
    append_jinja(find_paragraph(doc, "Participantes:"), " {{ participantes }}")

    # ── 1 map, under "Mapas:" ─────────────────────────────────────────────────
    imgs = image_paragraphs(doc)
    if len(imgs) != 1:
        raise RuntimeError(f"expected exactly 1 image paragraph in the mockup, found {len(imgs)}")
    replace_image_with_jinja(imgs[0], "{{ patrol_map }}")

    # ── 3 tables, in document order ───────────────────────────────────────────
    tables = doc.tables
    if len(tables) != 3:
        raise RuntimeError(f"expected exactly 3 tables in the mockup, found {len(tables)}")

    stats_tbl, events_tbl, photos_tbl = tables

    # Stats table: 6 fixed label/value rows (not a loop) — mockup's own row
    # order is ID da patrulha, Tipo de patrulha, Líder da patrulha,
    # Distância percorrida, Data e hora de início, Data e hora do fim.
    stat_exprs = [
        "{{ patrol_id }}",
        "{{ patrol_type }}",
        "{{ patrol_leader }}",
        "{{ total_distance_km }}",
        "{{ start_time }}",
        "{{ end_time }}",
    ]
    if len(stats_tbl.rows) != len(stat_exprs):
        raise RuntimeError(f"expected {len(stat_exprs)} rows in the stats table, found {len(stats_tbl.rows)}")
    for row, expr in zip(stats_tbl.rows, stat_exprs):
        set_cell_jinja(row.cells[1], expr)

    # Eventos table: 4th column is blank in the mockup itself (header and
    # sample row both blank) — left blank in every generated row too,
    # same treatment as the Relatório Bimestral template's Ameaças table.
    make_row_loop(
        events_tbl, 1, "events",
        [
            "{{ item['ID do evento'] }}",
            "{{ item['Tipo de evento'] }}",
            "{{ item['Detalhes do evento'] }}",
        ],
    )

    # Registros Fotográficos table.
    make_row_loop(
        photos_tbl, 1, "photos",
        ["{{ item.info }}", "{{ item.image }}"],
    )

    DST.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(DST))
    print(f"Template written to: {DST}")


if __name__ == "__main__":
    main()
